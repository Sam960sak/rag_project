import importlib.util
from pathlib import Path
import pytest

from src.parsers.document_parser import (
    DocumentParser, UnreadableDocumentError, UnsupportedFileTypeError)
from src.app.app import chunk_docs, STRATEGIES


def _make_txt(tmp_path):
    p = tmp_path / 'sample.txt'
    p.write_text('Machine Learning Basics\n\n'
                 'Supervised learning uses labeled datasets to train algorithms.\n\n'
                 'Unsupervised learning analyzes unlabeled data to find hidden patterns.\n',
                 encoding='utf-8')
    return p


def _make_docx(tmp_path):
    from docx import Document
    p = tmp_path / 'sample.docx'
    doc = Document()
    doc.add_heading('Vector Databases', level=1)
    doc.add_paragraph('Vector databases enable similarity search at scale.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Metric'
    table.cell(0, 1).text = 'Value'
    table.cell(1, 0).text = 'Recall@3'
    table.cell(1, 1).text = '0.90'
    doc.save(p)
    return p


def _make_pdf(tmp_path):
    from reportlab.pdfgen import canvas
    p = tmp_path / 'sample.pdf'
    c = canvas.Canvas(str(p))
    c.drawString(72, 720, 'Reinforcement learning involves agents learning')
    c.drawString(72, 700, 'through interaction with an environment.')
    c.showPage()
    c.save()
    return p


def _make_empty_pdf(tmp_path):
    from reportlab.pdfgen import canvas
    p = tmp_path / 'blank.pdf'
    c = canvas.Canvas(str(p))
    c.showPage()
    c.save()
    return p


def test_parse_txt(tmp_path):
    d = DocumentParser().parse_file(str(_make_txt(tmp_path)))
    assert 'Supervised learning' in d['text']
    assert d['pages'][0]['page'] == 1


def test_parse_docx_includes_table(tmp_path):
    d = DocumentParser().parse_file(str(_make_docx(tmp_path)))
    assert 'Vector databases enable similarity search' in d['text']
    assert 'Metric | Value' in d['text']
    assert 'Recall@3 | 0.90' in d['text']


def test_parse_pdf(tmp_path):
    d = DocumentParser().parse_file(str(_make_pdf(tmp_path)))
    assert 'Reinforcement learning' in d['text']
    assert d['metadata']['file_type'] == 'pdf'


def test_unreadable_pdf_raises_helpful_error(tmp_path):
    with pytest.raises(UnreadableDocumentError) as exc:
        DocumentParser().parse_file(str(_make_empty_pdf(tmp_path)))
    assert 'scanned' in str(exc.value).lower() or 'ocr' in str(exc.value).lower()


def test_unsupported_file_raises(tmp_path):
    p = tmp_path / 'data.csv'
    p.write_text('a,b,c', encoding='utf-8')
    with pytest.raises(UnsupportedFileTypeError):
        DocumentParser().parse_file(str(p))


def test_chunks_carry_source_and_page(tmp_path):
    docs = [DocumentParser().parse_file(str(_make_txt(tmp_path)))]
    ch = chunk_docs(docs)
    assert set(ch) == set(STRATEGIES)
    for name, chunks in ch.items():
        assert chunks, f'{name} produced no chunks'
        for c in chunks:
            assert c['metadata']['source'] == 'sample.txt'
            assert c['metadata']['page'] == 1
            assert 'PAGE' not in c['text']


@pytest.mark.skipif(importlib.util.find_spec('sentence_transformers') is None,
                    reason='sentence-transformers not installed')
def test_retrieval_all_strategies(tmp_path):
    from src.retrieval.retriever import Retriever
    docs = [DocumentParser().parse_file(str(_make_txt(tmp_path)))]
    ch = chunk_docs(docs)
    r = Retriever('all-MiniLM-L6-v2')  # model loaded once, reused for all strategies
    for name, chunks in ch.items():
        r.build(chunks)
        results = r.search('What is supervised learning?', top_k=3)
        assert results, f'{name} returned no results'
        assert all('similarity' in x and 'metadata' in x for x in results)
        # empty question must not crash and must return a list
        assert isinstance(r.search('', top_k=3), list)


@pytest.mark.skipif(importlib.util.find_spec('sentence_transformers') is None,
                    reason='sentence-transformers not installed')
def test_full_app_sample_mode_ui():
    from streamlit.testing.v1 import AppTest
    app_path = str(Path(__file__).resolve().parents[1] / 'run_app.py')
    at = AppTest.from_file(app_path, default_timeout=300)
    at.run()
    assert not at.exception

    # empty question -> helpful info, no crash
    at.button[0].click().run()
    assert not at.exception
    assert len(at.info) >= 1

    # real question -> retrieves across all three strategies without error
    at.text_input[0].set_value('What is supervised learning?').run()
    at.button[0].click().run()
    assert not at.exception
